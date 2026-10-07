"""Generate live pixel-only inspection reports in DOCX, PDF, and PNG formats."""

from datetime import datetime, timezone
import os
from textwrap import wrap

from backend.core.config import get_settings
from backend.schemas.report import ReportRequest


LIMITATION = "Measurements are currently reported in uncalibrated pixels. Physical dimensions require a validated scale reference and are not used in the current live assessment."
DISCLAIMER = "This system assesses visible surface damage from the provided image. It does not determine internal or subsurface damage and is not a structural-safety certification."


def _local_image(url: str | None) -> str | None:
    if not url or not url.startswith("/uploads/"):
        return None
    filename = os.path.basename(url[len("/uploads/"):])
    if not filename or filename != url[len("/uploads/"):]:
        return None
    path = os.path.abspath(os.path.join(get_settings().upload_dir, filename))
    root = os.path.abspath(get_settings().upload_dir)
    return path if os.path.commonpath([path, root]) == root and os.path.isfile(path) else None


def _safe_text(value, fallback="Not available") -> str:
    return fallback if value is None or value == "" else str(value)


def _rows(request: ReportRequest) -> list[list[str]]:
    rows = [["Detection", "Class", "Confidence", "Length (px)", "Width (px)", "Max width (px)", "Area (px^2)", "Orientation"]]
    for index, detection in enumerate(request.inference.detections, start=1):
        rows.append([
            str(index), detection.class_name, f"{detection.confidence * 100:.1f}%",
            _safe_text(detection.length_px), _safe_text(detection.width_px),
            _safe_text(detection.max_width_px), _safe_text(detection.area_px2),
            _safe_text(detection.orientation_deg),
        ])
    return rows


def _condition_lines(request: ReportRequest) -> list[str]:
    assessment = request.inference.condition_assessment
    if not assessment:
        return ["No condition assessment available."]
    return [
        f"Condition: {assessment.condition}",
        f"Condition score: {assessment.condition_score:.1f}",
        "Visible Surface Condition Assessment",
        "Measurement mode: Pixel / Uncalibrated",
        "Threshold profile: Project-defined prototype thresholds",
        f"Assessment basis: {', '.join(assessment.basis) if assessment.basis else 'No valid measurements used.'}",
        *[f"Reason: {reason}" for reason in assessment.reasons],
    ]


def _add_docx_image(document, url: str | None, caption: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    path = _local_image(url)
    if not path:
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(path, width=Inches(6.3))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _build_docx(request: ReportRequest, path: str, report_id: str, generated_at: str) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI-Based Wall Inspection System for Heritage Masonry")
    subtitle = document.add_paragraph("Automated Inspection Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Report identifier: {report_id}\nInspection date/time: {generated_at}")

    document.add_heading("Inspection Summary", level=1)
    summary = document.add_table(rows=0, cols=2)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in [("Image reference", _safe_text(request.image_reference)), ("Quality status", request.quality_status), ("Processing status", request.processing_status), ("Detection count", str(len(request.inference.detections)))]:
        cells = summary.add_row().cells
        cells[0].text, cells[1].text = label, value

    document.add_heading("Images", level=1)
    _add_docx_image(document, request.original_image_url, "Original image")
    _add_docx_image(document, request.processed_image_url, "Processed image")
    _add_docx_image(document, request.inference.overlay_image_url, "Detection overlay")

    document.add_heading("Damage Findings", level=1)
    findings = document.add_table(rows=0, cols=8)
    findings.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(_rows(request)):
        cells = findings.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    run.bold = row_index == 0

    document.add_heading("Summary Measurements", level=1)
    summary_measurements = request.inference.measurement_summary
    document.add_paragraph(f"Total detected cracks: {summary_measurements.detected_cracks if summary_measurements else 0}\nTotal crack length: {_safe_text(summary_measurements.total_length_px if summary_measurements else None)} px\nTotal crack area: {_safe_text(summary_measurements.total_area_px2 if summary_measurements else None)} px^2\nMaximum crack width: {_safe_text(summary_measurements.max_width_px if summary_measurements else None)} px")
    document.add_heading("Condition Assessment", level=1)
    for line in _condition_lines(request):
        document.add_paragraph(line)
    document.add_heading("Measurement Limitation", level=1)
    document.add_paragraph(LIMITATION)
    document.add_heading("Scope Disclaimer", level=1)
    document.add_paragraph(DISCLAIMER)
    document.save(path)


def _pdf_table(rows: list[list[str]]):
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import Table, TableStyle

    table = Table(rows, repeatRows=1, colWidths=[0.55 * inch, 0.7 * inch, 0.75 * inch, 0.72 * inch, 0.72 * inch, 0.82 * inch, 0.75 * inch, 0.82 * inch])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#34383b")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d9d9d9")), ("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f3f4")]), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    return table


def _build_pdf(request: ReportRequest, path: str, report_id: str, generated_at: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image as PdfImage, Paragraph, SimpleDocTemplate

    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, spaceAfter=6)
    heading = ParagraphStyle("Heading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=10, spaceAfter=6, textColor=colors.black)
    story = [Paragraph("AI-Based Wall Inspection System for Heritage Masonry", styles["Title"]), Paragraph("Automated Inspection Report", body), Paragraph(f"Report identifier: {report_id}<br/>Inspection date/time: {generated_at}", body), Paragraph("Inspection Summary", heading)]
    story.append(_pdf_table([["Field", "Value"], ["Image reference", _safe_text(request.image_reference)], ["Quality status", request.quality_status], ["Processing status", request.processing_status], ["Detection count", str(len(request.inference.detections))]]))
    story.append(Paragraph("Images", heading))
    for url, caption in [(request.original_image_url, "Original image"), (request.processed_image_url, "Processed image"), (request.inference.overlay_image_url, "Detection overlay")]:
        image_path = _local_image(url)
        if image_path:
            story.extend([PdfImage(image_path, width=6.4 * inch, height=3.2 * inch, kind="proportional"), Paragraph(caption, body)])
    story.extend([Paragraph("Damage Findings", heading), _pdf_table(_rows(request)), Paragraph("Summary Measurements", heading)])
    summary = request.inference.measurement_summary
    story.append(Paragraph(f"Total detected cracks: {summary.detected_cracks if summary else 0}<br/>Total crack length: {_safe_text(summary.total_length_px if summary else None)} px<br/>Total crack area: {_safe_text(summary.total_area_px2 if summary else None)} px^2<br/>Maximum crack width: {_safe_text(summary.max_width_px if summary else None)} px", body))
    story.append(Paragraph("Condition Assessment", heading)); story.extend(Paragraph(line, body) for line in _condition_lines(request))
    story.append(Paragraph("Measurement Limitation", heading)); story.append(Paragraph(LIMITATION, body))
    story.append(Paragraph("Scope Disclaimer", heading)); story.append(Paragraph(DISCLAIMER, body))
    SimpleDocTemplate(path, pagesize=letter, rightMargin=0.6 * inch, leftMargin=0.6 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch).build(story)


def _font(size: int):
    from PIL import ImageFont

    try:
        return ImageFont.truetype("arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def _build_png(request: ReportRequest, path: str, report_id: str, generated_at: str) -> None:
    from PIL import Image, ImageDraw

    canvas = Image.new("RGB", (1600, 1100), "#f5f3ef")
    draw = ImageDraw.Draw(canvas)
    title_font, heading_font, body_font = _font(38), _font(25), _font(20)
    draw.text((70, 55), "AI-Based Wall Inspection System for Heritage Masonry", fill="#1f2528", font=title_font)
    draw.text((70, 115), "Automated Inspection Report", fill="#8a5a45", font=heading_font)
    draw.text((70, 170), f"Report {report_id}  |  {generated_at}", fill="#60676b", font=body_font)
    draw.text((70, 240), "INSPECTION SUMMARY", fill="#1f2528", font=heading_font)
    summary = request.inference.measurement_summary
    lines = [f"Image: {_safe_text(request.image_reference)}", f"Quality: {request.quality_status}    Processing: {request.processing_status}", f"Detected cracks: {summary.detected_cracks if summary else 0}", f"Total length: {_safe_text(summary.total_length_px if summary else None)} px    Total area: {_safe_text(summary.total_area_px2 if summary else None)} px^2", f"Maximum width: {_safe_text(summary.max_width_px if summary else None)} px"]
    y = 285
    for line in lines:
        draw.text((70, y), line, fill="#30383d", font=body_font); y += 34
    draw.text((70, 485), "VISIBLE SURFACE CONDITION ASSESSMENT", fill="#1f2528", font=heading_font)
    assessment = request.inference.condition_assessment
    condition = assessment.condition if assessment else "No condition assessment available"
    score = f"Score: {assessment.condition_score:.1f}" if assessment else "Score: Not available"
    draw.text((70, 535), condition, fill="#8a5a45", font=_font(32)); draw.text((70, 580), score + "    Measurement mode: Pixel / Uncalibrated", fill="#30383d", font=body_font)
    overlay = _local_image(request.inference.overlay_image_url)
    if overlay:
        with Image.open(overlay) as source:
            source.thumbnail((560, 330))
            canvas.paste(source.convert("RGB"), (950, 230))
        draw.text((950, 575), "Detection overlay", fill="#60676b", font=body_font)
    draw.text((70, 790), "Project-defined prototype thresholds", fill="#30383d", font=body_font)
    for index, line in enumerate(wrap(LIMITATION, 105)):
        draw.text((70, 835 + index * 27), line, fill="#60676b", font=_font(17))
    for index, line in enumerate(wrap(DISCLAIMER, 105)):
        draw.text((70, 925 + index * 27), line, fill="#60676b", font=_font(17))
    canvas.save(path, "PNG")


def generate_report(request: ReportRequest) -> dict[str, str]:
    settings = get_settings()
    os.makedirs(settings.upload_dir, exist_ok=True)
    report_id = datetime.now(timezone.utc).strftime("AWIS-%Y%m%d-%H%M%S")
    generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    stem = f"{report_id}_inspection_report"
    docx_path = os.path.join(settings.upload_dir, stem + ".docx")
    pdf_path = os.path.join(settings.upload_dir, stem + ".pdf")
    png_path = os.path.join(settings.upload_dir, stem + ".png")
    _build_docx(request, docx_path, report_id, generated_at)
    _build_pdf(request, pdf_path, report_id, generated_at)
    _build_png(request, png_path, report_id, generated_at)
    return {"report_id": report_id, "docx_url": f"/uploads/{os.path.basename(docx_path)}", "pdf_url": f"/uploads/{os.path.basename(pdf_path)}", "png_url": f"/uploads/{os.path.basename(png_path)}"}
